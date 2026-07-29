#!/usr/bin/env python3
"""Assemble the Broken Wings manuscript markdown into one .docx."""
import re, glob, sys, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = sys.argv[2] if len(sys.argv) > 2 else "/home/user/schoolhouse-mobile/book/manuscript"
OUT = sys.argv[1] if len(sys.argv) > 1 else "/tmp/broken-wings.docx"

files = (["00-front-matter.md"]
         + sorted(os.path.relpath(f, BASE) for f in glob.glob(BASE + "/part-1/*.md"))
         + sorted(os.path.relpath(f, BASE) for f in glob.glob(BASE + "/part-2/*.md"))
         + sorted(os.path.relpath(f, BASE) for f in glob.glob(BASE + "/part-3/*.md"))
         + sorted(os.path.relpath(f, BASE) for f in glob.glob(BASE + "/part-4/*.md"))
         + [f for f in ("98-notes.md", "99-back-matter.md") if os.path.exists(f"{BASE}/{f}")])

doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Georgia"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, bold, before, after in [
    ("Heading 1", 20, True, 18, 12),
    ("Heading 2", 14, True, 14, 8),
]:
    st = doc.styles[name]
    st.font.name = "Georgia"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

part_style = doc.styles.add_style("PartTitle", WD_STYLE_TYPE.PARAGRAPH)
part_style.font.name = "Georgia"
part_style.font.size = Pt(26)
part_style.font.bold = True
part_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
part_style.paragraph_format.space_before = Pt(180)

quote_style = doc.styles.add_style("BookQuote", WD_STYLE_TYPE.PARAGRAPH)
quote_style.font.name = "Georgia"
quote_style.font.size = Pt(11)
quote_style.font.italic = True
quote_style.paragraph_format.left_indent = Inches(0.5)
quote_style.paragraph_format.space_after = Pt(6)

cap_style = doc.styles.add_style("FigCaption", WD_STYLE_TYPE.PARAGRAPH)
cap_style.font.name = "Georgia"
cap_style.font.size = Pt(9)
cap_style.font.italic = True
cap_style.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
cap_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap_style.paragraph_format.space_before = Pt(4)
cap_style.paragraph_format.space_after = Pt(14)

# margins
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.left_margin = sec.right_margin = Inches(1.1)

# page numbers in footer
def add_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for el, attrs, text in [("w:fldChar", {"w:fldCharType": "begin"}, None),
                            ("w:instrText", {"xml:space": "preserve"}, "PAGE"),
                            ("w:fldChar", {"w:fldCharType": "end"}, None)]:
        e = OxmlElement(el)
        for k, v in attrs.items():
            e.set(qn(k), v)
        if text:
            e.text = text
        run._r.append(e)

add_page_numbers(doc.sections[0])

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")

def add_runs(par, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(tok)

def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

def flush_para(buf, style=None):
    if not buf:
        return
    text = " ".join(l.strip() for l in buf).strip()
    if not text:
        buf.clear(); return
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text)
    buf.clear()

def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncols)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j in range(ncols):
            txt = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], txt)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


IMAGE = re.compile(r"^!\[(?P<cap>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")

fig_counter = {"n": 0}

def add_figure(caption, src):
    path = src if os.path.isabs(src) else os.path.join(BASE, src)
    if not os.path.exists(path):
        path = os.path.join(os.path.dirname(BASE.rstrip("/")), src)
    if not os.path.exists(path):
        print("  MISSING IMAGE", src)
        return
    fig_counter["n"] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(path, width=Inches(5.7))
    c = doc.add_paragraph(style="FigCaption")
    add_runs(c, f"Figure {fig_counter['n']}. {caption}")


# ---- pre-scan all top-level headings for a static Contents page ----
toc_entries = []
fig_entries = []
for fname in files:
    for line in open(f"{BASE}/{fname}", encoding="utf-8"):
        t = line.strip()
        if t.startswith("# ") and t[2:].strip() != "BROKEN WINGS":
            toc_entries.append(t[2:].strip())
        m = IMAGE.match(t)
        if m:
            fig_entries.append(m.group("cap"))

def emit_contents():
    page_break()
    h = doc.add_heading(level=1); h.add_run("Contents")
    for t in toc_entries:
        p = doc.add_paragraph()
        if t.startswith("PART "):
            r = p.add_run(t); r.bold = True
            p.paragraph_format.space_before = Pt(8)
        elif t.startswith("Chapter ") or t.startswith("Appendix") or t in (
            "Notes", "Epilogue — The Names", "A Note on Sources",
            "A Bibliographic Essay: Reading Further", "Acknowledgments", "Index"):
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, t)
        else:
            add_runs(p, t)
        p.paragraph_format.space_after = Pt(2)

    if fig_entries:
        page_break()
        h = doc.add_heading(level=1); h.add_run("List of Figures")
        for n, cap in enumerate(fig_entries, 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, f"Figure {n}. {cap}")
            p.paragraph_format.space_after = Pt(2)

first_heading = True
contents_done = False
for fname in files:
    lines = open(f"{BASE}/{fname}", encoding="utf-8").read().splitlines()
    buf, i = [], 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            flush_para(buf)
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(rows)
            continue
        m = IMAGE.match(s)
        if m:
            flush_para(buf)
            add_figure(m.group("cap"), m.group("src"))
            i += 1
            continue
        if s.startswith("# "):
            flush_para(buf)
            title = s[2:].strip()
            if not first_heading:
                if not contents_done:
                    emit_contents()
                    contents_done = True
                page_break()
            first_heading = False
            if title.startswith("PART "):
                doc.add_paragraph(title, style="PartTitle")
            elif title == "BROKEN WINGS":
                p = doc.add_paragraph(style="PartTitle")
                p.add_run(title)
            else:
                h = doc.add_heading(level=1)
                add_runs(h, title)
        elif s.startswith("## "):
            flush_para(buf)
            h = doc.add_heading(level=2)
            add_runs(h, s[3:].strip())
        elif s.startswith("> "):
            flush_para(buf)
            q = doc.add_paragraph(style="BookQuote")
            add_runs(q, s[2:].strip())
        elif s == ">":
            pass
        elif s == "---":
            flush_para(buf)
        elif s.startswith("- "):
            flush_para(buf)
            item = [s[2:]]
            while i + 1 < len(lines) and lines[i+1].startswith("  ") and lines[i+1].strip():
                i += 1; item.append(lines[i].strip())
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, " ".join(item))
        elif not s:
            flush_para(buf)
        else:
            buf.append(line)
        i += 1
    flush_para(buf)

doc.save(OUT)
print("saved", OUT)
